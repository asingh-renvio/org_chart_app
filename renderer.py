"""Excel-driven renderer that retains the approved organization-chart template."""
from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date
from html import escape
from io import BytesIO
from textwrap import wrap
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys

from openpyxl import load_workbook

REQUIRED_COLUMNS = {"Employee", "Department", "Manager", "Title"}
COLORS = {
    "pm": ("#534AB7", "#26215C", "#EEEDFE", "#DCD9FA", "#7F77DD"),
    "eng": ("#0F6E56", "#04342C", "#E1F5EE", "#C7EDDE", "#1D9E75"),
    "ai": ("#185FA5", "#042C53", "#E6F1FB", "#D0E4F7", "#378ADD"),
    "cs": ("#993C1D", "#4A1B0C", "#FAECE7", "#F6DDD2", "#D85A30"),
    "ga": ("#993556", "#4B1528", "#FBEAF0", "#F8DBE5", "#D4537E"),
    "sales": ("#854F0B", "#412402", "#FAEEDA", "#F8E0B5", "#BA7517"),
    "it": ("#5F5E5A", "#2C2C2A", "#E9E7DC", "#E9E7DC", "#888780"),
}
DEPARTMENTS = (
    ("CS", "Customer Support", "cs", ("Lynda Nicole Lee",)),
    ("GA", "General & Administrative", "ga", ("Shahla Ahsan", "Parveen Gulati")),
    ("RD", "Research & Development - Application Development", "eng", ("Giri Chandran",)),
    ("PM", "Product Management", "pm", ("Susheel Kumar", "Abhishek Jha", "Christina Logalbo")),
    ("AI", "Data Science & AI", "ai", ("Neeraj Sinha",)),
    ("SALES", "Sales & Marketing", "sales", ("Jeffrey Vizethann", "Saurabh Gupta")),
)

@dataclass
class Person:
    name: str
    title: str
    department: str
    manager: str
    chain: tuple[str, ...]
    inferred: bool = False


@dataclass
class Result:
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    people: list[Person] = field(default_factory=list)
    pdf: bytes | None = None
    png: bytes | None = None

    @property
    def valid(self) -> bool:
        return not self.errors


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _node_id(name: str, index: int) -> str:
    return f"N{index}_{re.sub(r'[^A-Za-z0-9_]', '_', name)}"


def _node(person: Person, node_id: str, department: str, big: bool = False) -> str:
    stroke, font, *_ = COLORS[department]
    title = '<br/>'.join(escape(line) for line in wrap(person.title, 28, break_long_words=False))
    title_row = f'<tr><td><font point-size="{12 if big else 10}">{title}</font></td></tr>' if title else ""
    size = 17 if big else 13
    return (f'{node_id} [label=<<table border="0" cellborder="0" cellspacing="0" cellpadding="2">'
            f'<tr><td><font point-size="{size}"><b>{escape(person.name)}</b></font></td></tr>{title_row}'
            f'</table>>, fillcolor="#FFFFFF", color="{stroke}", fontcolor="{font}"];')


def read_excel(upload: BytesIO) -> Result:
    result = Result()
    workbook = load_workbook(upload, data_only=True, read_only=True)
    sheet = workbook.active
    headers = {_text(value): index for index, value in enumerate(next(sheet.iter_rows(values_only=True)))}
    missing = sorted(REQUIRED_COLUMNS - set(headers))
    if missing:
        result.errors.append("Missing required columns: " + ", ".join(missing))
        return result
    status_index = headers.get("Employment status")
    active: dict[str, Person] = {}
    paths: list[tuple[str, ...]] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        name = _text(row[headers["Employee"]])
        if not name:
            continue
        if status_index is not None and _text(row[status_index]).casefold() != "active":
            continue
        key = name.casefold()
        if key in active:
            result.errors.append(f"Row {row_number}: duplicate active employee '{name}'.")
            continue
        chain = tuple(part.strip() for part in _text(row[headers["Manager"]]).split(">") if part.strip())
        active[key] = Person(name, _text(row[headers["Title"]]),
                     _text(row[headers["Department"]]) or "Unassigned",
                     chain[-1] if chain else "", chain)
        paths.append(chain)
    if not active:
        result.errors.append("No active employees found in the uploaded workbook.")
        return result
    inferred: dict[str, Person] = {}
    for employee in active.values():
        chain = employee.chain
        for index, name in enumerate(chain):
            key = name.casefold()
            if key not in active:
                inferred.setdefault(key, Person(name, "", employee.department,
                                                 chain[index - 1] if index else "",
                                                 chain[:index], True))
    result.people = list(active.values()) + list(inferred.values())
    if inferred:
        result.warnings.append(
            f"Added {len(inferred)} manager-only name(s) found in reporting paths: "
            + ", ".join(sorted(person.name for person in inferred.values())) + ".")
    _validate_reporting(result)
    return result


def _validate_reporting(result: Result) -> None:
    people = {person.name.casefold(): person for person in result.people}
    for person in result.people:
        if person.manager and person.manager.casefold() not in people:
            result.errors.append(f"'{person.name}' has an unknown direct manager '{person.manager}'.")
    for person in result.people:
        visited = {person.name.casefold()}
        current = person
        while current.manager:
            key = current.manager.casefold()
            if key in visited:
                result.errors.append(f"Reporting loop found involving '{person.name}'.")
                break
            visited.add(key)
            current = people.get(key, Person("", "", "", "", ()))
            if not current.name:
                break


def _department_for(person: Person) -> tuple[str, str, str]:
    path = person.department.casefold()
    names = {name.casefold() for name in person.chain + (person.name,)}
    if "customer support" in path:
        return "CS", "Customer Support", "cs"
    if "general & administrative" in path:
        return "GA", "General & Administrative", "ga"
    if "product management" in path:
        return "PM", "Product Management", "pm"
    if "sales" in path or "marketing" in path:
        return "SALES", "Sales & Marketing", "sales"
    if "it ops" in path:
        return "RD", "Research & Development - Application Development", "eng"
    if "neeraj sinha" in names or "data science" in person.title.casefold() or "ai engineer" in person.title.casefold():
        return "AI", "Data Science & AI", "ai"
    if "research & development" in path:
        return "RD", "Research & Development - Application Development", "eng"
    label = person.department.split(">")[0].strip() or "Unassigned"
    code = re.sub(r"[^A-Za-z0-9]", "_", label).strip("_").upper() or "UNASSIGNED"
    return code, label, "pm"


def _engineering_branch(person: Person) -> int:
    path = {name.casefold() for name in person.chain}
    if "ravikant kumar" in path or person.name.casefold() == "ravikant kumar":
        return 0
    if "prashant prashant" in path or person.name.casefold() == "prashant prashant":
        return 1
    return 2


def build_dot(people: list[Person]) -> tuple[str, list[str]]:
    """Build the approved department/team template from Excel reporting paths."""
    ids = {person.name.casefold(): _node_id(person.name, index) for index, person in enumerate(people, 1)}
    by_name = {person.name.casefold(): person for person in people}
    reports: dict[str, list[Person]] = defaultdict(list)
    for person in people:
        if person.manager:
            reports[person.manager.casefold()].append(person)
    unmatched: list[str] = []
    by_dept: dict[str, list[Person]] = defaultdict(list)
    department_meta: dict[str, tuple[str, str]] = {}
    for person in people:
        if not person.manager and person.inferred:
            continue
        code, label, color = _department_for(person)
        by_dept[code].append(person)
        department_meta[code] = (label, color)

    lines = ["digraph Organization {",
             'graph [rankdir=TB, newrank=true, splines=ortho, ranksep=0.6, nodesep=0.35, compound=true, '
             'label="Organization Chart", labelloc=t, fontsize=30, fontname="Helvetica-Bold", pad=0.6];',
             'node [shape=box, style="rounded,filled", fontname="Helvetica", penwidth=1.2, margin="0.14,0.08"];',
             'edge [color="#888780", penwidth=0.9, arrowsize=0.6];']
    heads = [person for person in people if not person.manager and person.inferred]
    if len(heads) != 1:
        raise ValueError("The uploaded workbook needs exactly one organization head in its Manager paths.")
    head = heads[0]
    lines.append(_node(head, ids[head.name.casefold()], "pm", big=True))
    anchor_ids: list[str] = []
    edge_lines: list[str] = []
    invisible: list[str] = []
    team_member_keys: set[str] = set()
    waypoint_seen: set[tuple[str, str]] = set()

    def append_team(lead: Person, children: list[Person], code: str,
                    color_key: str, team_number: int) -> None:
        stroke, font, background, team_background, team_stroke = COLORS[color_key]
        lines.append(f'subgraph cluster_{code}_{team_number} {{ label="{escape(lead.name)}\'s Team"; '
                     f'bgcolor="{team_background}"; color="{team_stroke}"; fontcolor="{font}"; fontsize=13; '
                     'fontname="Helvetica-Bold"; labeljust=c; style=rounded; margin=16;')
        lines.append(_node(lead, ids[lead.name.casefold()], color_key))
        lines.extend(_node(child, ids[child.name.casefold()], color_key) for child in children)
        for child_index, child in enumerate(children):
            team_member_keys.add(child.name.casefold())
            modifier = "" if child_index < 3 else " [constraint=false]"
            edge_lines.append(f'{ids[lead.name.casefold()]}->{ids[child.name.casefold()]}{modifier};')
            if child_index >= 3:
                invisible.append(f'{ids[children[child_index - 3].name.casefold()]}->'
                                 f'{ids[child.name.casefold()]} [style=invis];')
        lines.append("}")

    def append_team_rows(teams: list[tuple[Person, list[Person]]]) -> None:
        for row_start in range(0, len(teams), 3):
            row = teams[row_start:row_start + 3]
            if len(row) > 1:
                invisible.append("{rank=same; " + "; ".join(
                    ids[lead.name.casefold()] for lead, _ in row) + ";}")
        for team_index in range(3, len(teams)):
            previous_lead, previous_children = teams[team_index - 3]
            current_lead, _ = teams[team_index]
            previous_anchor = previous_children[-1] if previous_children else previous_lead
            invisible.append(f'{ids[previous_anchor.name.casefold()]}->{ids[current_lead.name.casefold()]} '
                             '[style=invis, weight=80];')

    preferred_order = [code for code, *_ in DEPARTMENTS]
    department_codes = [code for code in preferred_order if code in by_dept]
    department_codes.extend(sorted(set(by_dept) - set(department_codes)))
    for dept_index, code in enumerate(department_codes):
        label, color_key = department_meta[code]
        members = by_dept[code]
        if not members:
            continue
        stroke, font, background, team_background, team_stroke = COLORS[color_key]
        anchor = f"LAYOUT_{code}"
        anchor_ids.append(anchor)
        lines.append(f'subgraph cluster_{code} {{ label="{label}"; bgcolor="{background}"; color="{stroke}"; '
                     f'fontcolor="{font}"; fontsize=18; penwidth=1.6; fontname="Helvetica-Bold"; '
                     f'labeljust=c; style=rounded; margin=24; {anchor} [shape=point, width=0.01, height=0.01, label="", style=invis];')
        member_keys = {person.name.casefold() for person in members}
        team_leads = []
        team_children: dict[str, list[Person]] = {}
        for person in members:
            dept_reports = [child for child in reports[person.name.casefold()]
                             if child.name.casefold() in member_keys]
            leaf_children = [child for child in dept_reports if not reports[child.name.casefold()]]
            has_manager_reports = any(reports[child.name.casefold()] for child in dept_reports)
            # Senior managers who ALSO have sub-managers keep their loose individual
            # contributors as plain direct nodes instead of a self-titled team box
            # (matches the approved template's "Individual Contributors" grouping).
            if leaf_children and not has_manager_reports:
                team_leads.append(person)
                team_children[person.name.casefold()] = leaf_children
        team_keys = {person.name.casefold() for person in team_leads}
        child_keys = {child.name.casefold() for children in team_children.values() for child in children}
        direct = [person for person in members if person.name.casefold() not in team_keys
                  and person.name.casefold() not in child_keys]

        if code == "RD":
            # Each branch gets its own bordered box (dashed outline) so
            # Ravikant's and Prashant's organizations read as two distinct teams.
            branch_names = ("Ravikant Kumar", "Prashant Prashant")
            branch_member_keys: set[str] = set()
            for branch_index, branch_name in enumerate(branch_names, 1):
                branch = by_name.get(branch_name.casefold())
                if not branch:
                    continue
                branch_people = [person for person in members
                                 if _engineering_branch(person) == branch_index - 1]
                branch_member_keys.update(person.name.casefold() for person in branch_people)
                branch_leads = [lead for lead in team_leads if lead in branch_people]
                branch_direct = [person for person in direct
                                 if person in branch_people and person.name.casefold() != branch.name.casefold()]
                lines.append(f'subgraph cluster_RD_BRANCH_{branch_index} {{ label="{escape(branch.name)}\'s Organization"; '
                             f'bgcolor="{background}"; color="{team_stroke}"; fontcolor="{font}"; fontsize=14; '
                             'fontname="Helvetica-Bold"; labeljust=c; style="rounded,dashed"; margin=18;')
                lines.append(_node(branch, ids[branch.name.casefold()], color_key))
                lines.extend(_node(person, ids[person.name.casefold()], color_key) for person in branch_direct)
                branch_teams = [(lead, team_children[lead.name.casefold()]) for lead in branch_leads]
                for team_number, (lead, children) in enumerate(branch_teams, branch_index * 100):
                    append_team(lead, children, code, color_key, team_number)
                append_team_rows(branch_teams)
                lines.append("}")
            unbranched = [person for person in direct
                          if person.name.casefold() not in branch_member_keys]
            lines.extend(_node(person, ids[person.name.casefold()], color_key) for person in unbranched)
            ravikant = by_name.get("ravikant kumar")
            prashant = by_name.get("prashant prashant")
            if ravikant and prashant:
                invisible.append(f'{{rank=same; {ids[ravikant.name.casefold()]}; {ids[prashant.name.casefold()]};}}')
                invisible.append(f'{ids[ravikant.name.casefold()]}->{ids[prashant.name.casefold()]} '
                                 '[style=invis, weight=100];')
        else:
            lines.extend(_node(person, ids[person.name.casefold()], color_key) for person in direct)
            teams = [(lead, team_children[lead.name.casefold()]) for lead in team_leads]
            for team_number, (lead, children) in enumerate(teams, 1):
                append_team(lead, children, code, color_key, team_number)
            append_team_rows(teams)
        lines.append("}")
    for person in people:
        if not person.manager or person.manager.casefold() not in ids:
            continue
        if person.manager.casefold() == head.name.casefold():
            continue
        if person.name.casefold() not in team_member_keys:
            manager_person = by_name[person.manager.casefold()]
            person_code = _department_for(person)[0]
            same_department = person_code == _department_for(manager_person)[0]
            anchor = f"LAYOUT_{person_code}"
            if same_department:
                edge_lines.append(f'{ids[manager_person.name.casefold()]}->{ids[person.name.casefold()]};')
            elif anchor in anchor_ids:
                # Bend the cross-department line through the target department's
                # top anchor (sits above every box) instead of cutting straight
                # down through whichever box happens to be in the way.
                waypoint_key = (manager_person.name.casefold(), anchor)
                if waypoint_key not in waypoint_seen:
                    waypoint_seen.add(waypoint_key)
                    edge_lines.append(f'{ids[manager_person.name.casefold()]}->{anchor} '
                                     '[arrowhead=none, constraint=false, weight=0];')
                edge_lines.append(f'{anchor}->{ids[person.name.casefold()]} [constraint=false, weight=0];')
            else:
                edge_lines.append(f'{ids[manager_person.name.casefold()]}->{ids[person.name.casefold()]} '
                                 '[constraint=false, weight=0];')
    lines.extend(edge_lines)
    lines.extend(invisible)
    if anchor_ids:
        lines.append("{rank=same; " + "; ".join(anchor_ids) + ";}")
        for left, right in zip(anchor_ids, anchor_ids[1:]):
            lines.append(f"{left}->{right} [style=invis, weight=100];")
        for anchor in anchor_ids:
            weight = 100 if anchor == "LAYOUT_RD" else 0
            has_head_reports = any(f"LAYOUT_{_department_for(person)[0]}" == anchor
                                   for person in reports[head.name.casefold()])
            style = "solid" if has_head_reports else "invis"
            lines.append(f'{ids[head.name.casefold()]}->{anchor} '
                         f'[style={style}, arrowhead=none, weight={weight}];')
        roots_by_code = {code: roots for code, _, _, roots in DEPARTMENTS}
        for code in department_codes:
            if f"LAYOUT_{code}" not in anchor_ids:
                continue
            roots = roots_by_code.get(code, ())
            if not roots:
                roots = tuple(person.name for person in by_dept[code]
                              if not person.manager or person.manager.casefold() == head.name.casefold())
            head_reports = [person.name for person in by_dept[code]
                            if person.manager.casefold() == head.name.casefold()]
            roots = tuple(dict.fromkeys((*roots, *head_reports)))
            for root in roots:
                person = by_name.get(root.casefold())
                if person:
                    style = "solid" if person.manager.casefold() == head.name.casefold() else "invis"
                    lines.append(f"LAYOUT_{code}->{ids[person.name.casefold()]} [style={style}, weight=1];")
    lines.append("}")
    return "\n".join(lines), unmatched


def render(upload: BytesIO) -> Result:
    result = read_excel(upload)
    if not result.valid:
        return result
    try:
        dot, unmatched = build_dot(result.people)
    except ValueError as error:
        result.errors.append(str(error))
        return result
    if unmatched:
        result.errors.append("These people do not fit the approved department template: " + ", ".join(sorted(unmatched)))
        return result
    binary = Path(getattr(sys, "_MEIPASS", "")) / "graphviz" / "bin" / ("dot.exe" if os.name == "nt" else "dot")
    dot_binary = str(binary) if binary.is_file() else shutil.which("dot")
    if not dot_binary:
        result.errors.append("Graphviz is not installed on this server.")
        return result
    with subprocess.Popen([dot_binary, "-Tpdf"], stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE) as process:
        pdf, error = process.communicate(dot.encode())
        if process.returncode:
            result.errors.append("Graphviz PDF render failed: " + error.decode().strip())
            return result
    with subprocess.Popen([dot_binary, "-Tpng", "-Gdpi=144"], stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE) as process:
        png, error = process.communicate(dot.encode())
        if process.returncode:
            result.errors.append("Graphviz preview render failed: " + error.decode().strip())
            return result
    result.pdf, result.png = pdf, png
    return result